#!/usr/bin/env python3
"""Read-only audit of every supplied PLINK file; export aggregates only.

No participant IDs, individual genotypes, or raw source rows are exported.
Existing extracted files are read directly and compared with the archived CRCs.
The nested ZIP is opened only to read its directory; no source is extracted.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import time
import unicodedata
import zipfile
import zlib
from collections import Counter
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "04_데이터점검" / "2026-09-21"
TARGETS = {"rs429358", "rs7412"}
PUBLIC_SENTINELS = TARGETS | {"rs7903146", "rs671", "rs1229984", "rs9939609", "rs1801282"}
CHUNK = 4 * 1024 * 1024


def normal(value):
    return unicodedata.normalize("NFC", str(value))


def safe(value):
    if isinstance(value, dict):
        return {str(k): safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [safe(x) for x in value]
    if isinstance(value, np.generic):
        return value.item()
    return value


def snapshot(paths):
    return {str(p.relative_to(ROOT)): (p.stat().st_size, p.stat().st_mtime_ns) for p in paths}


def byte_digest(raw):
    return {"sha256": hashlib.sha256(raw).hexdigest(), "crc32": zlib.crc32(raw) & 0xFFFFFFFF,
            "bytes": len(raw)}


def read_codebook(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    sheets = []
    for sh in wb:
        sheets.append({"name": sh.title, "rows": sh.max_row, "columns": sh.max_column})
        for row_index, row in enumerate(sh.iter_rows(), 1):
            if any(cell.value is not None for cell in row):
                rows.append({"sheet": sh.title, "range": f"A{row_index}:{openpyxl.utils.get_column_letter(len(row))}{row_index}",
                             "values": [cell.value for cell in row]})
    wb.close()
    return {"path": str(path.relative_to(ROOT)), "sheets": sheets, "metadata_rows": rows}


def archive_directory(path):
    result = {"path": str(path.relative_to(ROOT)), "outer_crc_checked": False}
    meta = {}
    with zipfile.ZipFile(path) as outer:
        inner_name = next(i.filename for i in outer.infolist()
                          if normal(i.filename).endswith("유전체_합성데이터.zip"))
        # ZipExtFile supports seeking. This reads the archived member to the end,
        # validating its CRC, without retaining it or writing another raw file.
        with outer.open(inner_name) as stream:
            with zipfile.ZipFile(stream) as inner:
                inner_infos = inner.infolist()
                for i in inner_infos:
                    if not i.is_dir() and Path(i.filename).suffix in {".bed", ".bim", ".fam"}:
                        meta[Path(i.filename).name] = {"crc32": i.CRC, "bytes": i.file_size}
                result["inner_members"] = len(inner_infos)
                result["plink_members"] = len(meta)
                result["inner_directory_read"] = True
        result["inner_zip_outer_crc_checked_by_full_decompression"] = True
        checked = 1
        # Check every other outer member, including supplied codebooks and MRI ZIP.
        for i in outer.infolist():
            if i.filename == inner_name or i.is_dir():
                continue
            with outer.open(i) as stream:
                while stream.read(CHUNK):
                    pass
            checked += 1
        result["outer_crc_checked"] = True
        result["outer_members_crc_checked"] = checked
    return result, meta


def clinical_join(fam):
    by_id = {r[1]: r for r in fam}
    ids = set(by_id)
    result = {}
    paths = list((ROOT / "02_Data" / "cohort").glob("*.txt"))
    paths += list((ROOT / "02_Data").glob("*/*.csv"))
    for path in sorted(paths):
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            id_col = "T_ID" if "T_ID" in reader.fieldnames else "T00_ID"
            sex_col = "T_SEX" if "T_SEX" in reader.fieldnames else "T00_SEX" if "T00_SEX" in reader.fieldnames else None
            keys, sex_match, sex_mismatch = [], 0, 0
            for row in reader:
                key = row[id_col].strip()
                keys.append(key)
                if sex_col and key in by_id:
                    if row[sex_col].strip() == by_id[key][4]:
                        sex_match += 1
                    else:
                        sex_mismatch += 1
        other = set(keys)
        result[path.name] = {"id_column": id_col, "rows": len(keys), "unique_ids": len(other),
                             "matched_genome_ids": len(ids & other), "source_only_ids": len(other - ids),
                             "genome_only_ids": len(ids - other), "sex_column": sex_col,
                             "sex_matches": sex_match if sex_col else None,
                             "sex_mismatches": sex_mismatch if sex_col else None}
    return result


def genotype_block(path, index, n):
    size = (n + 3) // 4
    with path.open("rb") as f:
        f.seek(3 + index * size)
        raw = f.read(size)
    packed = np.frombuffer(raw, dtype=np.uint8)
    return ((packed[:, None] >> np.array([0, 2, 4, 6], dtype=np.uint8)) & 3).reshape(-1)[:n]


def main():
    started = time.time()
    OUT.mkdir(parents=True, exist_ok=True)
    source = next(p for p in (ROOT / "02_Data").iterdir() if p.is_dir() and normal(p.name) == "유전체_합성데이터")
    codebook = next(source.glob("*.xlsx"))
    archive = ROOT / "02_Data" / "Data.zip"
    sources = sorted(source.glob("*")) + [archive]
    before = snapshot(sources)
    report = {"audit_date": "2026-09-21", "privacy": "aggregate only; no participant IDs or individual genotypes exported",
              "source_directory": str(source.relative_to(ROOT)), "codebook": read_codebook(codebook),
              "executables": {"plink": shutil.which("plink"), "plink2": shutil.which("plink2")},
              "source_unchanged_check": "size and nanosecond mtime before and after; all sources opened read-only",
              "references": {"bed_format": "https://www.cog-genomics.org/plink/1.9/formats#bed",
                             "rs429358": "https://www.ncbi.nlm.nih.gov/snp/rs429358",
                             "rs7412": "https://www.ncbi.nlm.nih.gov/clinvar/variation/17848/"}}
    print("Checking archive directory and compressed-member CRCs", flush=True)
    report["archive"], archived = archive_directory(archive)
    reference_fam = None
    reference_raw = None
    seen_variants, seen_rs, leads, ld_parents = set(), set(), set(), set()
    all_types = Counter()
    file_manifest = []
    chromosomes = []
    target_records = {}
    public_metadata = []
    global_duplicate_extra_rows = 0
    cross_chr_duplicate_extra_rows = 0
    rs_duplicate_extra_rows = 0
    chromosome_rs = {}
    lead_chromosomes, ld_links = {}, []
    codes_per_byte = np.array([[(i >> shift) & 3 for shift in (0, 2, 4, 6)] for i in range(256)], dtype=np.uint8)
    code_lookup = np.array([(codes_per_byte == code).sum(axis=1) for code in range(4)], dtype=np.int64)
    for chrom in range(1, 23):
        prefix = source / f"synthetic_kchip_chr{chrom}"
        fam_path, bim_path, bed_path = [prefix.with_suffix(ext) for ext in (".fam", ".bim", ".bed")]
        fam_raw, bim_raw = fam_path.read_bytes(), bim_path.read_bytes()
        fam = [line.split() for line in fam_raw.decode().splitlines() if line.strip()]
        if any(len(row) != 6 for row in fam):
            raise ValueError(f"Malformed FAM on chromosome {chrom}; no source row disclosed")
        if reference_fam is None:
            reference_fam, reference_raw = fam, fam_raw
            id_counts = Counter(row[1] for row in fam)
            fid_counts = Counter(row[0] for row in fam)
            report["samples"] = {"n": len(fam), "unique_iid": len(id_counts), "unique_fid": len(fid_counts),
                                 "duplicate_iid_extra_rows": len(fam) - len(id_counts),
                                 "duplicate_fid_iid_extra_rows": len(fam) - len(set((r[0], r[1]) for r in fam)),
                                 "missing_or_zero_id_rows": sum(r[0] in {"", "0", "NA", "."} or r[1] in {"", "0", "NA", "."} for r in fam),
                                 "fid_equal_iid": sum(r[0] == r[1] for r in fam),
                                 "family_size_distribution": dict(Counter(fid_counts.values())),
                                 "paternal_id_nonzero": sum(r[2] != "0" for r in fam),
                                 "maternal_id_nonzero": sum(r[3] != "0" for r in fam),
                                 "sex_counts": dict(Counter(r[4] for r in fam)),
                                 "phenotype_counts": dict(Counter(r[5] for r in fam)),
                                 "missing_or_invalid_sex": sum(r[4] not in {"1", "2"} for r in fam),
                                 "missing_phenotype": sum(r[5] in {"-9", "0", "NA", "."} for r in fam),
                                 "family_caution": "Unique FID=IID and zero parents do not establish genetic unrelatedness"}
            report["id_linkage"] = clinical_join(fam)
        n = len(fam)
        local_ids, local_rs, local_positions, local_loci = set(), set(), set(), set()
        types, allele_pairs = Counter(), Counter()
        stat = {"chromosome": chrom, "samples": n, "fam_bytes_identical_to_chr1": fam_raw == reference_raw,
                "fam_all_columns_order_equal_chr1": fam == reference_fam,
                "fam_iid_order_equal_chr1": [r[1] for r in fam] == [r[1] for r in reference_fam],
                "variants": 0, "malformed_bim_rows": 0, "chromosome_label_mismatch": 0,
                "duplicate_id_extra_rows_within_chr": 0, "duplicate_position_extra_rows": 0,
                "duplicate_position_alleles_extra_rows": 0, "position_decreases": 0,
                "nonpositive_positions": 0, "nonzero_cm": 0, "invalid_dna_allele_rows": 0,
                "identical_alleles_rows": 0, "strand_ambiguous_rows": 0, "background_chr_label_mismatch": 0}
        previous = None
        positions = []
        for index, line in enumerate(bim_raw.decode().splitlines()):
            if not line.strip():
                continue
            fields = line.split()
            if len(fields) != 6:
                stat["malformed_bim_rows"] += 1
                continue
            label, variant, cm, bp, a1, a2 = fields
            pos = int(bp)
            stat["variants"] += 1
            stat["chromosome_label_mismatch"] += label != str(chrom)
            stat["nonpositive_positions"] += pos <= 0
            stat["nonzero_cm"] += float(cm) != 0
            stat["position_decreases"] += previous is not None and pos < previous
            previous = pos
            positions.append(pos)
            stat["duplicate_position_extra_rows"] += pos in local_positions
            local_positions.add(pos)
            locus = (pos, *sorted((a1, a2)))
            stat["duplicate_position_alleles_extra_rows"] += locus in local_loci
            local_loci.add(locus)
            if variant in seen_variants:
                global_duplicate_extra_rows += 1
                if variant not in local_ids:
                    cross_chr_duplicate_extra_rows += 1
            stat["duplicate_id_extra_rows_within_chr"] += variant in local_ids
            local_ids.add(variant)
            seen_variants.add(variant)
            if re.fullmatch(r"rs\d+", variant):
                typ = "rsid_lead"
                rs_duplicate_extra_rows += variant in seen_rs
                seen_rs.add(variant)
                local_rs.add(variant)
                leads.add(variant)
                lead_chromosomes[variant] = chrom
            elif re.fullmatch(r"rs\d+_ld\d+", variant):
                typ = "ld_neighbor"
                ld_parents.add(variant.rsplit("_ld", 1)[0])
                ld_links.append((chrom, variant.rsplit("_ld", 1)[0]))
            elif re.fullmatch(r"bg_chr\d+_\d+", variant):
                typ = "background"
                stat["background_chr_label_mismatch"] += int(re.match(r"bg_chr(\d+)_", variant).group(1)) != chrom
            else:
                typ = "other"
            types[typ] += 1
            stat["invalid_dna_allele_rows"] += a1 not in set("ACGT") or a2 not in set("ACGT")
            stat["identical_alleles_rows"] += a1 == a2
            stat["strand_ambiguous_rows"] += frozenset((a1, a2)) in {frozenset("AT"), frozenset("CG")}
            allele_pairs[a1 + "/" + a2] += 1
            if variant in PUBLIC_SENTINELS:
                public_metadata.append({"variant": variant, "chromosome": chrom, "position": pos, "A1": a1, "A2": a2})
            if variant in TARGETS:
                target_records[variant] = {"variant": variant, "chromosome": chrom, "position": pos, "A1": a1, "A2": a2,
                                           "index": index, "bed": bed_path, "n": n}
        stat.update({"min_bp": min(positions), "max_bp": max(positions), "rsid_leads": types["rsid_lead"],
                     "ld_neighbors": types["ld_neighbor"], "background_snps": types["background"], "other_ids": types["other"]})
        all_types.update(types)
        chromosome_rs[chrom] = local_rs
        h = hashlib.sha256()
        crc, byte_count = 0, 0
        hist = np.zeros(256, dtype=np.int64)
        with bed_path.open("rb") as f:
            magic = f.read(3)
            h.update(magic)
            crc = zlib.crc32(magic)
            byte_count = 3
            while block := f.read(CHUNK):
                h.update(block)
                crc = zlib.crc32(block, crc)
                byte_count += len(block)
                hist += np.bincount(np.frombuffer(block, dtype=np.uint8), minlength=256)
        expected = 3 + ((n + 3) // 4) * stat["variants"]
        stat.update({"bed_magic": magic.hex(), "bed_variant_major": magic == b"\x6c\x1b\x01",
                     "bed_bytes": byte_count, "bed_expected_bytes": expected, "bed_size_correct": byte_count == expected,
                     "all_genotype_missing_calls": int(code_lookup[1] @ hist) if n % 4 == 0 else None,
                     "all_genotype_call_count": n * stat["variants"]})
        for p, digest in ((fam_path, byte_digest(fam_raw)), (bim_path, byte_digest(bim_raw)),
                          (bed_path, {"sha256": h.hexdigest(), "crc32": crc & 0xFFFFFFFF, "bytes": byte_count})):
            expected_archive = archived.get(p.name)
            file_manifest.append({"path": str(p.relative_to(ROOT)), **digest,
                                  "matches_archived_crc_and_size": expected_archive == {"crc32": digest["crc32"], "bytes": digest["bytes"]}})
        stat["allele_pairs"] = dict(allele_pairs)
        chromosomes.append(stat)
        print(f"Audited chromosome {chrom}: {stat['variants']:,} variants; BED size {'OK' if stat['bed_size_correct'] else 'FAIL'}", flush=True)
    report["variants"] = {"total_rows": sum(x["variants"] for x in chromosomes), "unique_ids": len(seen_variants),
                          "id_types": dict(all_types), "unique_rsid_leads": len(leads),
                          "duplicate_id_extra_rows_global": global_duplicate_extra_rows,
                          "duplicate_id_extra_rows_cross_chr": cross_chr_duplicate_extra_rows,
                          "duplicate_rsid_extra_rows": rs_duplicate_extra_rows,
                          "ld_parent_leads_missing": len(ld_parents - leads),
                          "ld_parent_leads": len(ld_parents),
                          "ld_neighbors_parent_chr_mismatch": sum(lead_chromosomes.get(parent) != chrom for chrom, parent in ld_links),
                          "public_sentinel_variant_metadata": public_metadata}
    report["all_fam_identical"] = all(x["fam_bytes_identical_to_chr1"] for x in chromosomes)
    report["structural_issue_totals"] = {k: sum(c[k] for c in chromosomes) for k in (
        "malformed_bim_rows", "chromosome_label_mismatch", "duplicate_position_extra_rows",
        "duplicate_position_alleles_extra_rows", "position_decreases", "nonpositive_positions",
        "nonzero_cm", "invalid_dna_allele_rows", "identical_alleles_rows", "strand_ambiguous_rows",
        "background_chr_label_mismatch")}
    report["all_bed_magic_and_sizes_correct"] = all(x["bed_variant_major"] and x["bed_size_correct"] for x in chromosomes)
    report["all_local_plink_files_match_archive"] = all(x["matches_archived_crc_and_size"] for x in file_manifest)
    report["all_genotype_calls"] = sum(x["all_genotype_call_count"] for x in chromosomes)
    report["all_missing_genotype_calls"] = sum(x["all_genotype_missing_calls"] for x in chromosomes)
    report["genotype_scan_scope"] = "All BED bytes scanned for CRC/SHA and missing-code count; per-variant MAF/HWE, PCA, kinship and imputation not performed"
    report["genome_build"] = {"declared_in_codebook": False,
                               "caution": "Synthetic coordinates and allele labels must not be assumed to be real genome-build-aligned annotations"}
    apoe, joint = {}, {}
    blocks = {}
    for rsid in sorted(TARGETS):
        if rsid not in target_records:
            apoe[rsid] = {"present": False}
            continue
        record = target_records[rsid]
        codes = genotype_block(record["bed"], record["index"], record["n"])
        counts = np.bincount(codes, minlength=4)
        called = int(record["n"] - counts[1])
        a1, a2 = record["A1"], record["A2"]
        apoe[rsid] = {k: v for k, v in record.items() if k not in {"bed", "index", "n"}}
        apoe[rsid].update({"present": True, "samples": record["n"],
                          "genotype_counts": {a1 + "/" + a1: int(counts[0]), a1 + "/" + a2: int(counts[2]),
                                              a2 + "/" + a2: int(counts[3]), "missing": int(counts[1])},
                          "A1_frequency": float((2 * counts[0] + counts[2]) / (2 * called)),
                          "allele_pair_compatible_CT_or_complement_AG": set((a1, a2)) in ({"C", "T"}, {"A", "G"})})
        blocks[rsid] = codes
    if TARGETS <= set(blocks):
        c = np.bincount(blocks["rs429358"].astype(np.int64) * 4 + blocks["rs7412"], minlength=16).reshape(4, 4)
        joint = {"row_variant": "rs429358", "column_variant": "rs7412",
                 "code_order": ["A1/A1", "missing", "A1/A2", "A2/A2"], "counts": c.tolist()}
    report["APOE"] = {"variants": apoe, "joint_genotype_aggregate": joint,
                      "epsilon4_dosage_derivable_from_supplied_annotations": False,
                      "reason": "Both supplied APOE rsID records have non-APOE allele pairs and synthetic positions; rsID presence alone does not establish APOE haplotypes",
                      "reference_rs429358_positions": {"GRCh37": 45411941, "GRCh38": 44908684},
                      "reference_rs7412_positions": {"GRCh37": 45412079, "GRCh38": 44908822},
                      "safe_use": "Supplied synthetic candidate SNP dosage only; no epsilon2/3/4 labels without generator mapping confirmation"}
    report["research_feasibility"] = {
        "K_CHIP": "Supplied synthetic_kchip PLINK dataset; not evidence of genuine K-CHIP locus/allele/LD fidelity",
        "longitudinal_genomics": "No matched follow-up IDs; genetic prediction or association with within-person longitudinal MRI change cannot be evaluated in the supplied cohorts",
        "external_PRS": "Conventional externally weighted genome-wide PRS not established: 99.98% synthetic IDs, undeclared genome build and demonstrated APOE allele/coordinate mismatch. Even lead-only external weights require variant-by-variant harmonization or generator mapping",
        "internal_score": "A score fitted to synthetic SNP dosage is computationally feasible using training-only fitting; call it a synthetic-data genetic score rather than a validated real-world disease PRS",
        "GWAS": "Association scans can test the supplied simulation. Real-locus interpretation, gene/pathway mapping and biological causal claims require validated genomic annotations and generation assumptions",
        "APOE": "Do not compute or report epsilon4 carrier count from these two supplied rsID labels alone"}
    report["chromosomes"] = chromosomes
    report["file_manifest"] = file_manifest
    report["source_files_unchanged"] = snapshot(sources) == before
    report["elapsed_seconds"] = round(time.time() - started, 3)
    columns = [k for k in chromosomes[0] if k != "allele_pairs"]
    with (OUT / "genomics_chromosomes.tsv").open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(chromosomes)
    with (OUT / "genomics_summary.json").open("w", encoding="utf-8") as f:
        json.dump(safe(report), f, ensure_ascii=False, indent=2)
    print(json.dumps({"variants": report["variants"], "samples": report["samples"], "id_linkage": report["id_linkage"],
                      "APOE": report["APOE"], "all_files_match_archive": report["all_local_plink_files_match_archive"],
                      "all_missing_genotype_calls": report["all_missing_genotype_calls"],
                      "source_files_unchanged": report["source_files_unchanged"], "seconds": report["elapsed_seconds"]},
                     ensure_ascii=False, default=str, indent=2), flush=True)


if __name__ == "__main__":
    main()
