#!/usr/bin/env python3
"""Read-only MRI/codebook audit. Persist aggregates only, never person-level data.

Run with the bundled Python runtime from the project root. CSVs are read from
the delivered extracted folder. No source files are changed or re-exported.
"""
from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime
from itertools import combinations
from pathlib import Path
from zoneinfo import ZoneInfo

import numpy as np
import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "02_Data"
OUT = ROOT / "04_데이터점검" / "2026-09-21"
FILES = ["BrainVol_BASE", "BrainVol_FOLLOW_01", "BrainVol_FOLLOW_05"]
KEY_ROIS = ["BV_icv", "BV_subctx_hippocampus", "BV_whole_brain", "BV_ventricles"]
MISSING_TEXT = {"", "NA", "N/A", "NAN", "NULL", "NONE", "."}
SPECIAL = [66666, 77777, 88888, 99999]


def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def clean_json(x):
    if isinstance(x, dict):
        return {str(k): clean_json(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [clean_json(v) for v in x]
    if isinstance(x, (np.integer,)):
        return int(x)
    if isinstance(x, (float, np.floating)):
        return float(x) if np.isfinite(x) else None
    if isinstance(x, (np.bool_,)):
        return bool(x)
    return x


def stats(s):
    s = pd.Series(s, dtype=float).replace([np.inf, -np.inf], np.nan).dropna()
    if not len(s):
        return {"n": 0}
    q = s.quantile([.01, .05, .25, .5, .75, .95, .99])
    return {"n": len(s), "min": s.min(), "p01": q.loc[.01], "p05": q.loc[.05],
            "p25": q.loc[.25], "median": q.loc[.5], "p75": q.loc[.75],
            "p95": q.loc[.95], "p99": q.loc[.99], "max": s.max(),
            "mean": s.mean(), "sd": s.std(ddof=1)}


def clean_clinical_numeric(s):
    return pd.to_numeric(s, errors="coerce").replace(SPECIAL, np.nan)


def main():
    paths = {name: next(DATA.rglob(name + ".csv")) for name in FILES}
    source_dir = paths[FILES[0]].parent
    codebook_path = next(source_dir.glob("*.xlsx"))
    clinical_paths = {"BASE": DATA / "cohort" / "BASE_DATA1.txt",
                      "FOLLOW01": DATA / "cohort" / "FOLLOW_01_DATA.txt",
                      "FOLLOW05": DATA / "cohort" / "FOLLOW_05_DATA.txt"}
    sources = list(paths.values()) + [codebook_path] + list(clinical_paths.values())
    before = {str(p.relative_to(ROOT)): digest(p) for p in sources}
    wb = openpyxl.load_workbook(codebook_path, read_only=True, data_only=True)
    cb = {}
    for rownum, row in enumerate(wb["ROI_코드북"].iter_rows(values_only=True), 1):
        if len(row) > 1 and isinstance(row[1], str) and row[1].startswith("BV_"):
            cb[row[1]] = {"description_english": row[2], "description_korean": row[3],
                          "type": row[4], "unit": row[5], "codebook_sheet": "ROI_코드북",
                          "codebook_range": f"B{rownum}:F{rownum}"}
    overview = {"file": str(codebook_path.relative_to(ROOT)),
                "sheets": {ws.title: {"rows": ws.max_row, "columns": ws.max_column} for ws in wb},
                "overview_range": "개요!A4:D16", "roi_count": len(cb),
                "units": wb["개요"]["B13"].value,
                "nature": wb["개요"]["B14"].value,
                "declared_missing": wb["개요"]["B15"].value,
                "ROI_definition_note": wb["개요"]["A19"].value,
                "hippocampus_laterality": "One hippocampus volume column; no left/right hippocampus columns. The codebook does not explicitly label this value as a bilateral sum."}
    wb.close()
    report = {"audit_timestamp": datetime.now(ZoneInfo("Asia/Seoul")).isoformat(),
              "source_mode": "read delivered extracted CSVs without changing sources",
              "scope": {"all_rows_and_columns": True, "row_level_output": False,
                        "no_model_fitting": True, "no_cleaning_or_exclusion": True,
                        "extreme_rule": "Per-column 3*IQR and absolute robust MAD-z > 5 are review flags, not exclusion rules or medical limits."},
              "codebook": overview, "files": {}, "source_hashes_before": before}
    frames, numeric, rows = {}, {}, []
    for name, path in paths.items():
        with path.open(encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            header = next(reader)
            field_counts = {}
            for record in reader:
                field_counts[len(record)] = field_counts.get(len(record), 0) + 1
        df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        frames[name] = df
        idcol = df.columns[0]
        volumes = [c for c in df if c.startswith("BV_")]
        num = df[volumes].apply(pd.to_numeric, errors="coerce")
        numeric[name] = num
        info = {"file": str(path.relative_to(ROOT)), "rows": len(df), "columns": len(df.columns),
                "BV_columns": len(volumes), "id_column": idcol,
                "id_missing_n": int(df[idcol].str.strip().str.upper().isin(MISSING_TEXT).sum()),
                "id_unique_n": df[idcol].nunique(), "id_duplicate_rows": df[idcol].duplicated().sum(),
                "duplicate_complete_rows": df.duplicated().sum(), "duplicate_BV_profiles": num.duplicated().sum(),
                "duplicate_header_count": len(header) - len(set(header)), "row_field_counts": field_counts,
                "BV_columns_missing_from_codebook": sorted(set(volumes) - set(cb)),
                "codebook_BV_columns_missing_from_data": sorted(set(cb) - set(volumes)),
                "other_non_ID_columns": sorted(set(df) - set(volumes) - {idcol}),
                "BV_cells_scanned": int(num.size), "blank_or_text_missing_n": 0,
                "nonnumeric_n": 0, "nonfinite_n": 0, "nonpositive_n": 0,
                "special_token_n": 0, "extreme_review_cells_3IQR": 0,
                "key_ROIs": {c: stats(num[c]) for c in KEY_ROIS}}
        floor_columns = []
        for c in volumes:
            s = df[c].str.strip()
            val = num[c]
            miss = s.str.upper().isin(MISSING_TEXT)
            parse_fail = (~miss & val.isna())
            inf = val.notna() & ~np.isfinite(val)
            obs = val[np.isfinite(val)]
            q1, q3 = obs.quantile([.25, .75])
            lo, hi = q1 - 3 * (q3 - q1), q3 + 3 * (q3 - q1)
            mad = (obs - obs.median()).abs().median()
            extreme = obs.lt(lo) | obs.gt(hi)
            robust_n = int(((obs - obs.median()).abs() * .6744897501960817 / mad).gt(5).sum()) if mad else None
            r = {"file": path.name, "variable": c, **cb.get(c, {}), "rows": len(df),
                 "blank_or_text_missing_n": int(miss.sum()), "unexpected_nonnumeric_n": int(parse_fail.sum()),
                 "nonfinite_n": int(inf.sum()), "nonpositive_n": int(val.le(0).sum()),
                 "special_token_n": int(val.isin(SPECIAL).sum()),
                 "precision_not_exactly_3_decimals_n": int((~s.str.fullmatch(r"-?\d+\.\d{3}")).sum()),
                 "precision_more_than_3_decimals_n": int(s.str.contains(r"\.\d{4,}$", regex=True).sum()),
                 "distinct_values": int(obs.nunique()), "min_value_frequency": int(obs.eq(obs.min()).sum()),
                 "max_value_frequency": int(obs.eq(obs.max()).sum()), "review_lower_3IQR": lo,
                 "review_upper_3IQR": hi, "review_outside_3IQR_n": int(extreme.sum()),
                 "review_abs_robust_z_gt_5_n": robust_n, **stats(obs)}
            rows.append(r)
            info["blank_or_text_missing_n"] += r["blank_or_text_missing_n"]
            info["nonnumeric_n"] += r["unexpected_nonnumeric_n"]
            info["nonfinite_n"] += r["nonfinite_n"]
            info["nonpositive_n"] += r["nonpositive_n"]
            info["special_token_n"] += r["special_token_n"]
            info["extreme_review_cells_3IQR"] += r["review_outside_3IQR_n"]
            if r["min_value_frequency"] / len(obs) >= .01:
                floor_columns.append({"column": c, "minimum": obs.min(), "n_at_minimum": r["min_value_frequency"], "pct_at_minimum": 100*r["min_value_frequency"]/len(obs)})
        nonicv = num.drop(columns="BV_icv")
        info["ROI_gt_ICV_cells"] = int(nonicv.gt(num.BV_icv, axis=0).sum().sum())
        info["whole_brain_gt_ICV_n"] = int(num.BV_whole_brain.gt(num.BV_icv).sum())
        info["whole_brain_to_ICV_ratio"] = stats(num.BV_whole_brain / num.BV_icv)
        info["positive_whole_brain_minus_ICV_mL"] = stats((num.BV_whole_brain-num.BV_icv).where(num.BV_whole_brain.gt(num.BV_icv)))
        info["hippocampus_gt_whole_brain_n"] = int(num.BV_subctx_hippocampus.gt(num.BV_whole_brain).sum())
        info["ventricles_gt_ICV_n"] = int(num.BV_ventricles.gt(num.BV_icv).sum())
        info["repeated_minimum_at_least_1pct"] = floor_columns
        info["bilateral_sum_checks"] = {}
        for left in [c for c in volumes if "_lh_" in c]:
            right, total = left.replace("_lh_", "_rh_"), left.replace("_lh_", "_")
            if right in num and total in num:
                err = num[left] + num[right] - num[total]
                info["bilateral_sum_checks"][total] = {"left_column": left, "right_column": right,
                    "n_absolute_sum_difference_gt_0_002_mL": int(err.abs().gt(.002 + 1e-10).sum()),
                    "sum_minus_total_mL": stats(err), "left_right_identical_n": int(num[left].eq(num[right]).sum())}
        corr = num.corr()
        pairs = []
        for a, b in combinations(volumes, 2):
            if abs(corr.at[a, b]) >= .995:
                pairs.append({"column_a": a, "column_b": b, "pearson_r": corr.at[a,b]})
        info["absolute_correlation_at_least_0_995"] = pairs
        report["files"][name] = info

    report["id_intersections"] = {a+"__"+b: len(set(frames[a].iloc[:,0]) & set(frames[b].iloc[:,0])) for a,b in combinations(FILES,2)}
    clinical = {k: pd.read_csv(p, dtype=str, keep_default_na=False) for k,p in clinical_paths.items()}
    report["clinical_matching"] = {}
    for mri, cl, key in [(FILES[0], "BASE", "T_ID"), (FILES[1], "FOLLOW01", "T00_ID"), (FILES[2], "FOLLOW05", "T00_ID")]:
        a, b = set(frames[mri][key]), set(clinical[cl][key])
        report["clinical_matching"][mri] = {"MRI_n": len(a), "clinical_n": len(b), "matched_n": len(a&b),
            "MRI_only_n": len(a-b), "clinical_only_n": len(b-a),
            "same_ID_order": frames[mri][key].equals(clinical[cl][key])}
    f1 = frames[FILES[1]].set_index("T00_ID")
    f5 = frames[FILES[2]].set_index("T00_ID")
    common = f5.index.intersection(f1.index, sort=False)
    a = f1.loc[common].apply(pd.to_numeric, errors="coerce")
    b = f5.loc[common].apply(pd.to_numeric, errors="coerce")
    c1 = clinical["FOLLOW01"].set_index("T00_ID").reindex(common)
    c5 = clinical["FOLLOW05"].set_index("T00_ID").reindex(common)
    d1 = pd.to_datetime(c1.T01_EDATE, format="%Y%m", errors="coerce")
    d5 = pd.to_datetime(c5.T05_EDATE, format="%Y%m", errors="coerce")
    months = (d5.dt.year-d1.dt.year)*12 + d5.dt.month-d1.dt.month
    years = months / 12
    report["longitudinal"] = {"FOLLOW01_n": len(f1), "FOLLOW05_n": len(f5), "paired_n": len(common),
        "FOLLOW05_not_in_FOLLOW01_n": len(f5.index.difference(f1.index)),
        "FOLLOW01_not_in_FOLLOW05_n": len(f1.index.difference(f5.index)),
        "two_delivered_MRI_visits_only": True, "missing_intermediate_visits": ["02", "03", "04"],
        "elapsed_time_source": "clinical T01_EDATE and T05_EDATE in YYYYMM; month difference / 12, not exact scan-date timestamps",
        "elapsed_years": stats(years), "missing_elapsed_years_n": int(years.isna().sum()),
        "nonpositive_elapsed_years_n": int(years.le(0).sum()),
        "annual_change_computable_n": int((years.gt(0) & a[KEY_ROIS].notna().all(axis=1) & b[KEY_ROIS].notna().all(axis=1)).sum()),
        "symmetric_annual_percent_formula": "200*(V05-V01)/(V05+V01)/elapsed_years",
        "positive_loss_formula": "-1 * symmetric_annual_percent_change; higher is faster volume loss",
        "key_ROI_changes": {}}
    for c in a:
        delta = b[c]-a[c]
        annual = 200*delta/(a[c]+b[c])/years
        # Keep per-ROI longitudinal aggregate diagnostics alongside file columns.
        row = next(r for r in rows if r["file"] == "BrainVol_FOLLOW_05.csv" and r["variable"] == c)
        row.update({"paired_n": len(common), "paired_decrease_n": int(delta.lt(0).sum()),
                    "paired_unchanged_n": int(delta.eq(0).sum()), "paired_increase_n": int(delta.gt(0).sum()),
                    "paired_annual_symmetric_pct_mean": annual.mean(), "paired_annual_symmetric_pct_sd": annual.std(),
                    "paired_annual_symmetric_pct_min": annual.min(), "paired_annual_symmetric_pct_max": annual.max()})
        if c in KEY_ROIS:
            r = {"delta_mL": stats(delta), "annual_symmetric_percent_change": stats(annual),
                 "decrease_n": int(delta.lt(0).sum()), "unchanged_n": int(delta.eq(0).sum()), "increase_n": int(delta.gt(0).sum()),
                 "baseline_followup_pearson_r": a[c].corr(b[c])}
            if c != "BV_icv":
                aa, bb = a[c]/a.BV_icv, b[c]/b.BV_icv
                annual_current_icv = 200*(bb-aa)/(bb+aa)/years
                base_icv_a, base_icv_b = a[c]/a.BV_icv, b[c]/a.BV_icv
                annual_fixed_icv = 200*(base_icv_b-base_icv_a)/(base_icv_b+base_icv_a)/years
                r["annual_visit_specific_ICV_normalized_percent_change"] = stats(annual_current_icv)
                r["sign_flips_using_visit_specific_ICV_n"] = int((np.sign(annual_current_icv) != np.sign(annual)).sum())
                r["fixed_baseline_ICV_matches_raw_symmetric_change_max_abs_error"] = (annual_fixed_icv-annual).abs().max()
            report["longitudinal"]["key_ROI_changes"][c] = r
    icv_pct = 200*(b.BV_icv-a.BV_icv)/(b.BV_icv+a.BV_icv)
    report["longitudinal"]["ICV_symmetric_change_pct"] = stats(icv_pct)
    report["longitudinal"]["ICV_absolute_symmetric_change_gt_5pct_n"] = int(icv_pct.abs().gt(5).sum())
    report["longitudinal"]["ICV_absolute_symmetric_change_gt_10pct_n"] = int(icv_pct.abs().gt(10).sum())
    report["longitudinal"]["ventricle_unchanged_values"] = {}
    for c in ["BV_ventricles", "BV_lateral_ventricle_w_inf", "BV_lateral_ventricle", "BV_inf_lat_vent"]:
        same = a[c].eq(b[c])
        report["longitudinal"]["ventricle_unchanged_values"][c] = {
            "unchanged_n": int(same.sum()), "unchanged_at_baseline_minimum_n": int((same & a[c].eq(a[c].min())).sum()),
            "decrease_n": int(b[c].lt(a[c]).sum()), "increase_n": int(b[c].gt(a[c]).sum())}
    report["synthetic_pattern_review"] = [
        "BASE and FOLLOW01 have repeated exact lower limits in four ventricular variables; this suggests a generation boundary but its mechanism is not known.",
        "Many longitudinal ventricular values are exactly unchanged despite nonzero elapsed time; the unchanged values are not simply observations at the minimum baseline volume.",
        "Paired left/right volume sums differ from their named total far beyond 0.001 mL rounding in most rows; preserve supplied aggregate ROI variables rather than recomputing them from hemispheres.",
        "Some BASE whole-brain volumes exceed ICV. Flag for source clarification and sensitivity analysis; no values were corrected or excluded.",
        "Statistics describe the supplied synthetic data only and do not validate biological realism or clinical utility."]
    report["longitudinal"]["interpretation"] = [
        "Two volume observations permit change-rate calculation, not individual nonlinear trajectories.",
        "ICV varies between synthetic visits; use baseline-fixed ICV for normalized longitudinal comparisons and report sensitivity to visit-specific ICV.",
        "A percent decline is negative under the defined formula; define faster loss as its negative or select the bottom 20% of signed change.",
        "Clinical examination months are available for all pairs, but MRI acquisition dates are not included."]
    report["research_plan_v4_1_review"] = [
        {"status": "supported", "topic": "10,000 cross-sectional and 1,000-to-698 follow-up datasets; each has 66 BV variables; ICV, hippocampus, whole brain and ventricles available"},
        {"status": "revise", "topic": "Concept diagram still says follow-up 5 timepoints; actual delivery contains 01 and 05 only"},
        {"status": "revise", "topic": "Hippocampus is one available volume variable; bilateral-sum interpretation is not explicit in the codebook"},
        {"status": "revise", "topic": "Clarify signed annual change versus positive loss: upper 20% signed change does not mean fastest atrophy"},
        {"status": "revise", "topic": "Section 7 score-by-time test remains inconsistent with a two-visit change-rate regression plan"},
        {"status": "clarify", "topic": "A small cross-sectional age/sex adjusted volume is a research phenotype, not proof of acquired atrophy or a clinically validated abnormality threshold"},
        {"status": "clarify", "topic": "Longitudinal denominator is examination-month interval, unless imaging dates are supplied; account for changing synthetic ICV"},
        {"status": "clarify", "topic": "Synthetic value distributions, floors and volume-sum relationships require model sensitivity checks, not automatic source corrections"}]
    after = {str(p.relative_to(ROOT)): digest(p) for p in sources}
    report["source_integrity"] = {"all_source_hashes_unchanged": before == after, "sha256": after}
    report.pop("source_hashes_before")
    report["aggregate_columns_rows"] = len(rows)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "mri_summary.json").write_text(json.dumps(clean_json(report), ensure_ascii=False, indent=2, allow_nan=False)+"\n", encoding="utf-8")
    pd.DataFrame(rows).to_csv(OUT / "mri_columns.tsv", sep="\t", index=False)
    # This console output includes no individual-level keys, records or values.
    print(json.dumps(clean_json({"files": {k: {c:v[c] for c in ["rows", "columns", "BV_columns", "id_duplicate_rows", "blank_or_text_missing_n", "nonnumeric_n", "nonpositive_n"]} for k,v in report["files"].items()},
          "id_intersections": report["id_intersections"], "clinical_matching": report["clinical_matching"],
          "paired_n": len(common), "elapsed_years": stats(years), "source_integrity_ok": before == after}), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
