#!/usr/bin/env python3
"""Read every clinical record locally; write aggregate-only audit results.

No individual identifiers, records, or joined individual-level data are exported.
Original TXT/XLSX files are opened read-only. No statistical model is fitted.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
from itertools import combinations
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "02_Data" / "cohort"
OUT = ROOT / "04_데이터점검" / "2026-09-21"
SPECIAL = {"66666": "not_surveyed", "77777": "not_applicable", "99999": "unknown_or_unmeasured"}
MISSING_TEXT = {"", ".", "NA", "N/A", "NULL", "NAN"}
FILES = ["BASE_DATA1", "BASE_DATA2", "BASE_DATA3", "FOLLOW_01_DATA", "FOLLOW_05_DATA"]


def safe(value):
    if isinstance(value, dict):
        return {str(k): safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [safe(x) for x in value]
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not np.isfinite(value):
        return None
    return value


def num_summary(series):
    s = pd.to_numeric(series, errors="coerce").dropna()
    if not len(s):
        return {"n": 0}
    q = s.quantile([0, .01, .25, .5, .75, .99, 1])
    return {"n": len(s), "mean": float(s.mean()), "sd": float(s.std()),
            **{k: float(v) for k, v in zip(["min", "p01", "p25", "median", "p75", "p99", "max"], q)}}


def clean_num(series):
    text = series.astype(str).str.strip()
    n = pd.to_numeric(text, errors="coerce")
    return n.mask(n.isin([int(x) for x in SPECIAL]))


def miss_mask(series):
    text = series.astype(str).str.strip()
    return text.str.upper().isin(MISSING_TEXT) | pd.to_numeric(text, errors="coerce").isin([int(x) for x in SPECIAL])


def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def canonical(col):
    key = re.sub(r"^T(?:00|01|05)?_", "", col)
    return {"CREATINE": "CREATININE", "FMFHT1": "FMFHT", "FMFHT2": "FMMHT",
            "FMFDM1": "FMFDM", "FMFDM2": "FMMDM", "PMYN_C": "PMYN", "PMAG_C": "PMAG"}.get(key, key)


def codebook_read():
    books, maps = {}, {}
    for path in sorted(SOURCE.glob("*.xlsx")):
        normalized = unicodedata.normalize("NFC", path.name)
        key = "BASE" if normalized.startswith("[기반조사]") else "FOLLOW05" if "4차추적" in normalized else "FOLLOW01"
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        records = {}
        sheet_metadata = []
        for sh in wb:
            sheet_metadata.append({"name": sh.title, "rows": sh.max_row, "columns": sh.max_column})
            for row in sh.iter_rows(min_row=4):
                if len(row) < 7 or not row[3].value:
                    continue
                name = str(row[3].value).strip()
                records[name] = {"description": str(row[4].value or "").strip(),
                                 "codes_or_units": str(row[5].value or "").strip(),
                                 "type": str(row[6].value or "").strip(),
                                 "questionnaire": str(row[7].value or "").strip() if len(row) > 7 else "",
                                 "sheet": sh.title, "range": f"D{row[3].row}:G{row[3].row}",
                                 "codebook": str(path.relative_to(ROOT))}
        books[key] = {"file": str(path.relative_to(ROOT)), "sheets": sheet_metadata, "variables": len(records)}
        maps[key] = records
        wb.close()
    return books, maps


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [SOURCE / f"{name}.txt" for name in FILES]
    paths += sorted(SOURCE.glob("*.xlsx"))
    hashes_before = {str(p.relative_to(ROOT)): digest(p) for p in paths}
    books, codebooks = codebook_read()
    frames, rows, report = {}, [], {"audit_date": "2026-09-21", "privacy": "aggregate-only; no individual IDs or records exported",
        "scope": {"clinical_text_files": 5, "clinical_codebooks": len(books),
                  "checks": ["every CSV row field count", "all IDs for uniqueness, missingness and set intersections",
                             "all cells for blanks, missing tokens, special codes and numeric parsing",
                             "all codebook categorical fields for allowed-code membership",
                             "every observed numeric column summarized after special-code exclusion",
                             "all 698 matched follow-up intervals and age/date consistency",
                             "all available height/weight/BMI, BP-order and diagnosis-age checks",
                             "all columns mapped to the clinical codebooks; cross-cohort common-variable names",
                             "retained vs non-retained baseline aggregate comparisons"],
                  "not_done": ["No raw-value correction, imputation, subject exclusion, prediction or phenotype construction",
                               "No claim to validate measurement accuracy or synthetic biological realism",
                               "No MRI/genotype audit in this script", "No external lookup or network access"]},
        "codebooks": books, "files": {}}
    for name in FILES:
        path = SOURCE / f"{name}.txt"
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            header = next(reader)
            field_counts = {}
            for line in reader:
                field_counts[len(line)] = field_counts.get(len(line), 0) + 1
        df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        frames[name] = df
        key = "BASE" if name.startswith("BASE") else "FOLLOW01" if "01" in name else "FOLLOW05"
        codebook = codebooks[key]
        id_col = df.columns[0]
        info = {"rows": len(df), "columns": len(df.columns), "header_duplicate_count": len(header) - len(set(header)),
                "row_field_counts": field_counts, "id_column": id_col,
                "id_missing": int(miss_mask(df[id_col]).sum()), "id_unique": int(df[id_col].nunique()),
                "id_duplicate_rows": int(df[id_col].duplicated().sum()), "all_row_duplicate_count": int(df.duplicated().sum()),
                "data_class_counts": df.iloc[:, 1].value_counts().to_dict(),
                "columns_not_in_codebook": sorted(set(df.columns) - set(codebook)),
                "empty_cell_count": int(df.eq("").sum().sum()), "special_counts": {},
                "columns_with_invalid_codes": {}, "columns_with_unexpected_nonnumeric": {}}
        class_col = df.columns[1]
        allowed_classes = set(re.findall(r"[BF]\d{2}", codebook.get(class_col, {}).get("codes_or_units", "")))
        info["data_class_invalid_n"] = int((~df[class_col].isin(allowed_classes)).sum()) if allowed_classes else None
        for special in SPECIAL:
            info["special_counts"][special] = sum(int(pd.to_numeric(df[c], errors="coerce").eq(int(special)).sum()) for c in df)
        for c in df:
            s = df[c].str.strip()
            cb = codebook.get(c, {})
            idlike = c.endswith("_ID") or c.endswith("DATA_CLASS")
            numeric = clean_num(s)
            missing = miss_mask(s)
            record = {"file": path.name, "variable": c, "canonical": canonical(c), "rows": len(df),
                      "blank_or_text_missing": int(s.str.upper().isin(MISSING_TEXT).sum()),
                      **{f"n_{special}": int(pd.to_numeric(s, errors="coerce").eq(int(special)).sum()) for special in SPECIAL},
                      "missing_or_special": int(missing.sum()), "missing_or_special_pct": float(missing.mean() * 100),
                      "observed_valid": int((~missing).sum()), "distinct_observed": int(s[~missing].nunique()),
                      **cb}
            if not idlike:
                nonnumeric = int((~missing & numeric.isna()).sum())
                record["unexpected_nonnumeric"] = nonnumeric
                if nonnumeric:
                    info["columns_with_unexpected_nonnumeric"][c] = nonnumeric
                record.update(num_summary(numeric))
                if "범주" in cb.get("type", ""):
                    allowed = {int(x) for x in re.findall(r"(?<!\d)(\d+)\s*=", cb.get("codes_or_units", ""))}
                    if allowed:
                        invalid = numeric.notna() & ~numeric.isin(allowed)
                        record["invalid_category_count"] = int(invalid.sum())
                        record["allowed_codes"] = ",".join(map(str, sorted(allowed)))
                        record["observed_code_counts"] = json.dumps(numeric.dropna().value_counts().sort_index().to_dict())
                        if invalid.any():
                            info["columns_with_invalid_codes"][c] = {"n": int(invalid.sum()), "code_counts": numeric[invalid].value_counts().to_dict()}
            rows.append(record)
        report["files"][name] = info

    report["id_overlap"] = {a + "__" + b: len(set(frames[a].iloc[:, 0]) & set(frames[b].iloc[:, 0])) for a, b in combinations(FILES, 2)}
    base1, base2, base3 = [frames[f"BASE_DATA{i}"] for i in (1, 2, 3)]
    base = base1.merge(base2.drop(columns="T_DATA_CLASS"), on="T_ID", validate="one_to_one").merge(base3.drop(columns="T_DATA_CLASS"), on="T_ID", validate="one_to_one")
    report["base_join"] = {"required_operation": "horizontal 1:1 join on T_ID; not vertical concatenation",
        "rows": len(base), "columns": len(base.columns),
        "same_id_order": all(base1["T_ID"].equals(x["T_ID"]) for x in [base2, base3]),
        "data_class_mismatches_after_id_join": {f"BASE_DATA{i}": int((base1.set_index("T_ID")["T_DATA_CLASS"] != frames[f"BASE_DATA{i}"].set_index("T_ID")["T_DATA_CLASS"].reindex(base1["T_ID"])).sum()) for i in (2, 3)},
        "codebook_variables_missing_from_join": sorted(set(codebooks["BASE"]) - set(base.columns)),
        "extra_variables_not_in_codebook": sorted(set(base.columns) - set(codebooks["BASE"]))}
    f1, f5 = frames["FOLLOW_01_DATA"], frames["FOLLOW_05_DATA"]
    paired = f1.merge(f5, on="T00_ID", suffixes=("_01", "_05"), validate="one_to_one")
    report["longitudinal"] = {"baseline_n": len(f1), "last_visit_n": len(f5), "matched_n": len(paired),
        "not_in_last_visit_n": len(f1) - len(paired), "retained_pct": len(paired) / len(f1) * 100,
        "follow05_ids_not_in_follow01": len(set(f5.T00_ID) - set(f1.T00_ID)),
        "baseline10k_and_follow01_overlap": len(set(base.T_ID) & set(f1.T00_ID)),
        "delivered_clinical_visits": ["01", "05"],
        "intermediate_clinical_files_missing": ["FOLLOW_02_DATA.txt", "FOLLOW_03_DATA.txt", "FOLLOW_04_DATA.txt"],
        "sex_comparison_limit": "FOLLOW_05 has no sex field; use T00_SEX from FOLLOW_01, cannot independently test sex stability",
        "data_class_mismatch": int((paired.T00_DATA_CLASS_01 != paired.T00_DATA_CLASS_05).sum())}
    report["longitudinal"]["disease_semantics"] = {
        "baseline": "Diagnosis history during lifetime",
        "follow05": "Diagnosis since the previous cohort examination, not cumulative prevalent disease",
        "evidence": {"codebook": codebooks["FOLLOW05"]["T05_HTN"]["codebook"],
                     "sheet": codebooks["FOLLOW05"]["T05_HTN"]["sheet"], "cell": "H27",
                     "questionnaire": codebooks["FOLLOW05"]["T05_HTN"]["questionnaire"]},
        "consequence": "Do not interpret baseline yes/follow05 no as recovery or an error. With intermediate clinical visits absent, total 01-to-05 cumulative incidence cannot be reconstructed directly."}
    date_months = {}
    report["dates"] = {}
    for name, col in [("BASE_DATA1", "T_EDATE"), ("FOLLOW_01_DATA", "T01_EDATE"), ("FOLLOW_05_DATA", "T05_EDATE")]:
        s = frames[name][col]
        parsed = pd.to_datetime(s, format="%Y%m", errors="coerce")
        report["dates"][name] = {"column": col, "resolution": "month (YYYYMM)", "invalid_n": int(parsed.isna().sum()),
            "length_counts": s.str.len().value_counts().to_dict(), "min": parsed.min().strftime("%Y-%m"), "max": parsed.max().strftime("%Y-%m")}
    p1 = pd.to_datetime(paired.T01_EDATE, format="%Y%m", errors="coerce")
    p5 = pd.to_datetime(paired.T05_EDATE, format="%Y%m", errors="coerce")
    months = (p5.dt.year - p1.dt.year) * 12 + (p5.dt.month - p1.dt.month)
    years = months / 12
    age_change = clean_num(paired.T05_AGE) - clean_num(paired.T01_AGE)
    report["longitudinal"].update({"elapsed_months": num_summary(months), "elapsed_years": num_summary(years),
        "nonpositive_intervals_n": int(months.le(0).sum()), "age_increase": num_summary(age_change),
        "age_change_minus_elapsed_years": num_summary(age_change - years),
        "abs_age_change_date_discrepancy_gt_1_year": int((age_change - years).abs().gt(1).sum()),
        "age_decrease_n": int(age_change.lt(0).sum())})

    maps = {"BASE": {canonical(c): c for c in base.columns}, "FOLLOW01": {canonical(c): c for c in f1.columns}, "FOLLOW05": {canonical(c): c for c in f5.columns}}
    common = set(maps["BASE"]) & set(maps["FOLLOW01"])
    report["variables"] = {"base_follow01_common_canonical": sorted(common),
        "base_only_canonical_vs_follow01": sorted(set(maps["BASE"]) - set(maps["FOLLOW01"])),
        "follow01_only_canonical_vs_base": sorted(set(maps["FOLLOW01"]) - set(maps["BASE"])),
        "follow01_not_in_follow05_canonical": sorted(set(maps["FOLLOW01"]) - set(maps["FOLLOW05"])),
        "common_name_caution": "Canonical name harmonization is metadata only; categories and source questions still require semantic harmonization.",
        "derived_BMI_available": "FOLLOW01/FOLLOW05 BMI can be derived in memory from kg and cm when both are observed; baseline has T_BMI.",
        "LDL": "T_LDL only in BASE; no direct LDL in FOLLOW01 or FOLLOW05.",
        "education": "T01_EDU only in FOLLOW01; no education field in BASE.",
        "income": "T_INCOME only in BASE; no income field in FOLLOW01/FOLLOW05."}

    report["logic_checks"] = {}
    report["cohort_profile"] = {}
    report["range_flags"] = {}
    # Broad plausibility screens identify values for review, not automatic deletion rules.
    bounds = {"HEIGHT": (100, 230), "WEIGHT": (20, 250), "BMI": (10, 80), "WAIST": (40, 200), "HIP": (40, 220),
              "SBP": (60, 280), "DBP": (30, 180), "PULSE": (25, 220), "HBA1C": (3, 25), "GLU0": (20, 700),
              "CREATININE": (.1, 20), "AST": (0, 3000), "ALT": (0, 3000), "TCHL": (50, 700), "HDL": (5, 200),
              "LDL": (0, 600), "TG": (10, 4000), "INS0": (0, 500)}
    for label, df in [("BASE", base), ("FOLLOW01", f1), ("FOLLOW05", f5)]:
        mp = maps[label]
        n = {k: clean_num(df[c]) for k, c in mp.items() if k not in ["ID", "DATA_CLASS"]}
        sex = n.get("SEX")
        if sex is None:
            sex = df.T00_ID.map(f1.set_index("T00_ID").T00_SEX).astype(float)
        profile = {"age": num_summary(n["AGE"]), "sex_counts": sex.value_counts().to_dict()}
        checks = {"SBP_below_DBP": int((n["SBP"] < n["DBP"]).sum()), "disease_diagnosis_age_after_current_age": {}}
        for disease in ["HTN", "DM", "LIP"]:
            profile[disease + "_counts"] = n[disease].value_counts().to_dict()
            checks["disease_diagnosis_age_after_current_age"][disease] = int((n[disease + "AG"] > n["AGE"]).sum())
        bmi = n["WEIGHT"] / (n["HEIGHT"] / 100) ** 2
        profile["BMI_from_height_weight"] = num_summary(bmi)
        if "BMI" in n:
            diff = n["BMI"] - bmi
            checks["BMI_provided_minus_calculated"] = num_summary(diff)
            checks["BMI_absolute_difference_gt_0_1"] = int(diff.abs().gt(.1).sum())
            checks["BMI_absolute_difference_gt_0_5"] = int(diff.abs().gt(.5).sum())
            checks["BMI_provided_equals_rounded_calculated_n"] = int((n["BMI"] == bmi.round()).sum())
        checks["female_specific_observations_in_males"] = {key: int((sex.eq(1) & n[key].notna()).sum()) for key in ["MNSAG", "PMYN", "PMAG", "PREG", "FPREGAG", "CHILD", "FLABAG"] if key in n}
        report["logic_checks"][label] = checks
        report["cohort_profile"][label] = profile
        report["range_flags"][label] = {key: {"review_lower": lo, "review_upper": hi,
            "below_n": int(n[key].lt(lo).sum()), "above_n": int(n[key].gt(hi).sum())} for key, (lo, hi) in bounds.items() if key in n}

    retained = f1.T00_ID.isin(f5.T00_ID)
    attrition = {"interpretation": "Absent from delivered final-visit file; reason is not recorded and must not be assumed to be death or dropout."}
    for group_name, mask in [("retained", retained), ("not_in_follow05", ~retained)]:
        subset = f1.loc[mask]
        d = {"n": len(subset), "age": num_summary(clean_num(subset.T01_AGE)), "sex_counts": subset.T00_SEX.value_counts().to_dict()}
        for c in ["T01_HTN", "T01_DM", "T01_LIP", "T01_SMOKE", "T01_DRINK", "T01_EXER"]:
            val = clean_num(subset[c])
            d[c] = {"missing_n": int(val.isna().sum()), "counts": val.value_counts().to_dict()}
        attrition[group_name] = d
    attrition["retention_by_data_class"] = []
    for klass, sub in f1.groupby("T00_DATA_CLASS"):
        nret = int(sub.T00_ID.isin(f5.T00_ID).sum())
        attrition["retention_by_data_class"].append({"class": klass, "baseline_n": len(sub), "retained_n": nret, "retained_pct": nret / len(sub) * 100})
    report["attrition"] = attrition
    report["longitudinal"]["disease_transition_counts"] = {}
    report["longitudinal"]["disease_transition_counts_caution"] = "The two visits ask different time windows; these cross-tabulations are not cumulative disease-state transitions."
    for disease in ["HTN", "DM", "LIP"]:
        a, b = clean_num(paired[f"T01_{disease}"]), clean_num(paired[f"T05_{disease}"])
        report["longitudinal"]["disease_transition_counts"][disease] = {"no_to_no": int((a.eq(1) & b.eq(1)).sum()),
            "no_to_yes": int((a.eq(1) & b.eq(2)).sum()), "yes_to_yes": int((a.eq(2) & b.eq(2)).sum()),
            "yes_to_no": int((a.eq(2) & b.eq(1)).sum()), "any_missing": int((a.isna() | b.isna()).sum())}

    # Completeness of planned common clinical model blocks; no imputation is done.
    c1 = ["AGE", "SEX", "SBP", "DBP", "WAIST", "HEIGHT", "WEIGHT", "SMOKE", "DRINK", "EXER", "HTN", "DM", "LIP", "FMFHT", "FMMHT", "FMFDM", "FMMDM"]
    c2 = ["GLU0", "HBA1C", "TCHL", "HDL", "TG", "AST", "ALT", "CREATININE"]
    report["model_input_completeness"] = {"C1_definition": c1, "C2_additions": c2, "note": "These complete-case counts are an audit, not a recommendation to discard incomplete participants."}
    for label, df in [("BASE", base), ("FOLLOW01", f1)]:
        mp = maps[label]
        report["model_input_completeness"][label] = {}
        for block, names in [("C1", c1), ("C1_C2", c1+c2), ("C1_C2_without_HBA1C", c1+[k for k in c2 if k != "HBA1C"])]:
            columns = [mp[k] for k in names]
            any_missing = pd.concat([miss_mask(df[c]) for c in columns], axis=1).any(axis=1)
            report["model_input_completeness"][label][block] = {"n": len(df), "complete_n": int((~any_missing).sum()), "any_missing_n": int(any_missing.sum())}

    report["hba1c_missing_by_data_class"] = {}
    for label, df, class_col, val_col in [("BASE", base, "T_DATA_CLASS", "T_HBA1C"), ("FOLLOW01", f1, "T00_DATA_CLASS", "T01_HBA1C"), ("FOLLOW05", f5, "T00_DATA_CLASS", "T05_HBA1C")]:
        group_rows = []
        for klass, sub in df.groupby(class_col):
            group_rows.append({"class": klass, "n": len(sub), "missing_n": int(miss_mask(sub[val_col]).sum()),
                               "not_surveyed_66666_n": int(pd.to_numeric(sub[val_col], errors="coerce").eq(66666).sum())})
        report["hba1c_missing_by_data_class"][label] = group_rows
    report["semantics_notes"] = [
        "Data class is a source/unit-data label, not a physiological predictor; examine source/cohort and missingness shifts.",
        "Baseline T_MARRY codes 3/4/5 mean separated/divorced/widowed; FOLLOW01 and FOLLOW05 codes 3/4/5 mean widowed/separated/divorced.",
        "DRINK numeric codes map to never/former/current in both cohorts even though yes/no wording is reversed by the source questionnaire.",
        "T_TOTALC and T_PACKYR are present only in BASE_DATA2 and not documented in the supplied baseline codebook; names alone do not establish units/derivation.",
        "Broad range flags are review prompts, not exclusion or diagnostic rules; measurement accuracy is not validated."]

    hashes_after = {str(p.relative_to(ROOT)): digest(p) for p in paths}
    report["source_integrity"] = {"all_source_hashes_unchanged": hashes_before == hashes_after, "sha256": hashes_after}
    report["output_files"] = [str((OUT / "clinical_summary.json").relative_to(ROOT)), str((OUT / "clinical_columns.tsv").relative_to(ROOT))]
    (OUT / "clinical_summary.json").write_text(json.dumps(safe(report), ensure_ascii=False, indent=2, allow_nan=False)+"\n", encoding="utf-8")
    pd.DataFrame(rows).to_csv(OUT / "clinical_columns.tsv", sep="\t", index=False)
    print(json.dumps(safe({"files": {k: {kk: v[kk] for kk in ["rows", "columns", "id_duplicate_rows", "columns_not_in_codebook", "columns_with_invalid_codes"]} for k, v in report["files"].items()},
        "base_join": report["base_join"], "longitudinal": report["longitudinal"], "logic_checks": report["logic_checks"],
        "model_input_completeness": report["model_input_completeness"], "source_integrity_ok": hashes_before == hashes_after}), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
